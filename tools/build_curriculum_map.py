#!/usr/bin/env python3
"""Bygg kursplanekartan för världshistoria som Excelfil åt Akelius (Daniel Shea).

Läser data/curriculum-map-history.json (kartan: nivåmål, ämnesrader per nivågrupp,
standardval) och data/cards.json (de befintliga korten) och skriver
export/Akelius-history-curriculum-map-<datum>.xlsx med flikarna
Read me, Summary, Topics, Level outcomes, Standards och Existing cards.
Summeringarna är formler, så de räknas om när någon ändrar antalet kort.

Kontrollerar före skrivning att alla 80 befintliga set ligger i kartan med
1 + 1 + 3 kort, att summan befintliga kort är 400 och att varje målkod som
en ämnesrad hänvisar till finns.

Användning:
  python3 tools/build_curriculum_map.py [--date 2026-10-15]
"""
import json, pathlib, re, sys
from collections import Counter, defaultdict
from datetime import date

ROOT = pathlib.Path(__file__).resolve().parent.parent
MAP = ROOT / 'data' / 'curriculum-map-history.json'
OUT = ROOT / 'export'

TIERS = ['1-3', '4-6', '7-9', '10-U']
LEVELS = [str(n) for n in range(1, 13)] + ['U']
TIER_OF_LEVEL = {**{str(n): '1-3' for n in (1, 2, 3)}, **{str(n): '4-6' for n in (4, 5, 6)},
                 **{str(n): '7-9' for n in (7, 8, 9)}, **{str(n): '10-U' for n in (10, 11, 12)}, 'U': '10-U'}
LEVEL_OF_CARD = {'grade 6': '6', 'grade 9': '9', 'grade 12': '12', 'university': 'U'}


def tier_label(t):
    # "Tier 1-3" i stället för "1-3": Excel tolkar annars kriteriet "1-3" som ett datum i COUNTIF
    return f'Tier {t}'


def check(m, cards):
    problems = []
    slugs = [s['slug'] for s in cards['sets']]
    placed = defaultdict(Counter)
    for r in m['topics']:
        for slug in [x.strip() for x in r['existing_slug'].split(',') if x.strip()]:
            placed[slug][r['tier']] += 1
    for slug in slugs:
        if slug not in placed:
            problems.append(f'befintligt set saknas i kartan: {slug}')
            continue
        for t in ('4-6', '7-9', '10-U'):
            if placed[slug][t] != 1:
                problems.append(f'{slug}: {placed[slug][t]} rader i nivågrupp {t}, väntat 1')
    unknown = set(placed) - set(slugs)
    if unknown:
        problems.append(f'okända set i kartan: {sorted(unknown)}')
    total_existing = sum(r['existing_cards'] for r in m['topics'])
    if total_existing != 400:
        problems.append(f'befintliga kort summerar till {total_existing}, väntat 400')
    ids = {o['outcome_id'] for o in m['level_outcomes']}
    for r in m['topics']:
        for oid in re.split(r'[,;]\s*', r['outcome_ids'].strip()):
            if oid and oid not in ids:
                problems.append(f"{r['topic_id']}: okänd målkod {oid}")
    per_level = Counter(o['level'] for o in m['level_outcomes'])
    for lv in LEVELS:
        if not 4 <= per_level[lv] <= 8:
            problems.append(f'nivå {lv}: {per_level[lv]} mål, väntat 4-8')
    for field in ('approach_summary', 'totals_note'):
        if '—' in m.get(field, ''):
            problems.append(f'tankstreck i {field}')
    return problems


def build(m, cards, path, when):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    F = 'Arial'
    base, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    head_fill = PatternFill('solid', start_color='1F3864')
    head_font = Font(name=F, size=10, bold=True, color='FFFFFF')
    wrap = Alignment(wrap_text=True, vertical='top')
    thin = Side(style='thin', color='BFBFBF')
    wb = Workbook()

    def sheet(title, cols, widths, rows):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for row in rows:
            ws.append(row)
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font, cell.alignment = base, wrap
        ws.freeze_panes = 'A2'
        ref = f'A1:{get_column_letter(len(cols))}{ws.max_row}'
        tab = Table(displayName=re.sub(r'\W', '', title), ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(tab)
        return ws

    # --- Read me ---------------------------------------------------------------------------
    rm = wb.active
    rm.title = 'Read me'
    sc = m['standard_choice']
    n_new = sum(r['new_cards'] for r in m['topics'])
    lines = [
        ('Akelius Quiz Cards · History curriculum map', 'title'),
        (f'Proposal from Jonas von Essen, {when}. Draft for discussion with Daniel Shea.', None),
        ('', None),
        ('What this file is', 'h'),
        ('A model curriculum for the history cards from level 1 to university, as asked for in the email '
         '"Quiz Card curricula" of 15 October 2026: key topics per tier, a rough number of cards per topic, '
         'outcomes for every level, and a mapping to an academic standard.', None),
        ('Levels are school years 1 to 12 plus U for university. A card is labelled by subject and level, '
         'for example H 3. Tiers group the levels: 1-3, 4-6, 7-9 and 10-U.', None),
        ('400 cards already exist: 80 topics with five cards each, at levels 6, 9, 12, 12 and U. '
         f'The map places all of them and proposes {n_new} new cards.', None),
        ('', None),
        ('Sheets', 'h'),
        ('Summary: cards per tier, region and strand. The numbers are formulas and update when the Topics sheet changes.', None),
        ('Topics: the curriculum map. One row per topic and tier, with levels, content, card numbers, outcomes and standard mapping.', None),
        ('Level outcomes: what students can do by the end of each level, with the standard each outcome maps to.', None),
        ('Standards: the documents the map is built on, with links, and the ones it does not use.', None),
        ('Existing cards: the 400 cards that are written, with their level and the map row they belong to.', None),
        ('', None),
        ('How to edit', 'h'),
        ('Change card numbers in the New cards column of the Topics sheet; totals update by themselves. '
         'Add a topic by adding a row to the Topics table. Outcome ids link the two sheets, for example H8.3.', None),
        ('', None),
        ('Choice of standard', 'h'),
        (f'Primary standard: {sc["primary"]}', 'b'),
        ('Supporting: ' + '; '.join(sc['supporting']), None),
        (sc['rationale'], None),
        (sc['not_used'], None),
        ('', None),
        ('How the map is organised', 'h'),
        (m['approach_summary'], None),
        ('', None),
        ('Card numbers', 'h'),
        (m['totals_note'], None),
    ]
    for r, (txt, kind) in enumerate(lines, 1):
        c = rm.cell(row=r, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.font = (Font(name=F, size=14, bold=True) if kind == 'title' else
                  Font(name=F, size=11, bold=True) if kind == 'h' else
                  bold if kind == 'b' else base)
    rm.column_dimensions['A'].width = 120

    # --- Topics ----------------------------------------------------------------------------
    topics = sorted(m['topics'], key=lambda r: (TIERS.index(r['tier']), r['topic_id']))
    trow = ['Topic ID', 'Tier', 'Levels', 'Topic', 'Key content', 'Era', 'Region', 'Strand', 'Status',
            'Existing cards', 'New cards', 'Total cards', 'Outcomes', 'Standard mapping',
            'Existing card set', 'Notes']
    rows = []
    for i, r in enumerate(topics, 2):
        rows.append([r['topic_id'], tier_label(r['tier']), r['levels'], r['topic'], r['key_content'], r['era'],
                     r['region'], r['strand'], r['status'], r['existing_cards'], r['new_cards'],
                     f'=J{i}+K{i}', r['outcome_ids'], r['standard_refs'], r['existing_slug'], r['notes']])
    ws_t = sheet('Topics', trow, [11, 10, 9, 30, 50, 22, 16, 22, 13, 9, 9, 9, 18, 45, 26, 40], rows)
    last_t = ws_t.max_row
    yellow = PatternFill('solid', start_color='FFF2CC')
    for i in range(2, last_t + 1):
        ws_t[f'K{i}'].fill = yellow

    # --- Level outcomes --------------------------------------------------------------------
    serves = defaultdict(list)
    for r in topics:
        for oid in re.split(r'[,;]\s*', r['outcome_ids'].strip()):
            if oid:
                serves[oid].append(r['topic_id'])
    outs = sorted(m['level_outcomes'], key=lambda o: (LEVELS.index(o['level']),
                                                      [int(x) if x.isdigit() else x for x in re.split(r'(\d+)', o['outcome_id'])]))
    orows = [[o['level'], tier_label(TIER_OF_LEVEL[o['level']]), o['outcome_id'], o['statement'], o['standard_refs'],
              ', '.join(serves.get(o['outcome_id'], []))] for o in outs]
    sheet('Level outcomes', ['Level', 'Tier', 'Outcome ID', 'By the end of the level, students can', 'Standard mapping',
                             'Topic rows that teach it'], [7, 10, 10, 60, 45, 40], orows)

    # --- Standards -------------------------------------------------------------------------
    srows = [[s['role'], s['name'], s['code_version'], s['levels'], s['used_for'], s['url']] for s in m['standards_list']]
    sheet('Standards', ['Role', 'Standard', 'Code and version', 'Levels', 'Used for', 'Link'], [14, 38, 30, 10, 55, 60], srows)

    # --- Existing cards --------------------------------------------------------------------
    row_of = {}
    for r in topics:
        for slug in [x.strip() for x in r['existing_slug'].split(',') if x.strip()]:
            row_of[(slug, r['tier'])] = r['topic_id']
    erows = []
    for n, s in enumerate(cards['sets'], 1):
        for k, c in enumerate(s['cards'], 1):
            lv = LEVEL_OF_CARD[c['level']]
            t = TIER_OF_LEVEL[lv]
            erows.append([f'{n:03d}-{k}', s['person'].replace(' - ', ': '), k, f'H {lv}', tier_label(t), c['title'],
                          c['question'], c['correct'], row_of.get((s['slug'], t), ''), s['slug']])
    sheet('Existing cards', ['Card ID', 'Topic', 'Card', 'Level', 'Tier', 'Title', 'Question', 'Answer', 'Map row',
                             'Card set'], [9, 30, 6, 7, 10, 40, 30, 26, 11, 26], erows)

    # --- Summary (formler) -----------------------------------------------------------------
    su = wb.create_sheet('Summary', 1)
    su.column_dimensions['A'].width = 34
    for col in 'BCDEF':
        su.column_dimensions[col].width = 14
    r0 = 1
    su.cell(row=r0, column=1, value='Cards per tier').font = Font(name=F, size=12, bold=True)
    hdr = ['Tier', 'Topic rows', 'Existing cards', 'New cards', 'Total cards', 'Share of total']
    for c, h in enumerate(hdr, 1):
        cell = su.cell(row=r0 + 1, column=c, value=h)
        cell.font, cell.fill = head_font, head_fill
    T = f'Topics!$B$2:$B${last_t}'
    first = r0 + 2
    for j, t in enumerate(TIERS):
        rr = first + j
        su.cell(row=rr, column=1, value=tier_label(t))
        su.cell(row=rr, column=2, value=f'=COUNTIF({T},A{rr})')
        su.cell(row=rr, column=3, value=f'=SUMIF({T},A{rr},Topics!$J$2:$J${last_t})')
        su.cell(row=rr, column=4, value=f'=SUMIF({T},A{rr},Topics!$K$2:$K${last_t})')
        su.cell(row=rr, column=5, value=f'=C{rr}+D{rr}')
    tot = first + len(TIERS)
    su.cell(row=tot, column=1, value='All tiers').font = bold
    for col in 'BCDE':
        su[f'{col}{tot}'] = f'=SUM({col}{first}:{col}{tot - 1})'
        su[f'{col}{tot}'].font = bold
    for rr in range(first, tot + 1):
        su[f'F{rr}'] = f'=IF($E${tot}=0,0,E{rr}/$E${tot})'
        su[f'F{rr}'].number_format = '0%'

    def breakdown(start, title, col_letter, values):
        su.cell(row=start, column=1, value=title).font = Font(name=F, size=12, bold=True)
        for c, h in enumerate(['Value', 'Topic rows', 'Existing cards', 'New cards', 'Total cards', 'Share of total'], 1):
            cell = su.cell(row=start + 1, column=c, value=h)
            cell.font, cell.fill = head_font, head_fill
        rng = f'Topics!${col_letter}$2:${col_letter}${last_t}'
        for j, v in enumerate(values):
            rr = start + 2 + j
            su.cell(row=rr, column=1, value=v)
            su.cell(row=rr, column=2, value=f'=COUNTIF({rng},A{rr})')
            su.cell(row=rr, column=3, value=f'=SUMIF({rng},A{rr},Topics!$J$2:$J${last_t})')
            su.cell(row=rr, column=4, value=f'=SUMIF({rng},A{rr},Topics!$K$2:$K${last_t})')
            su.cell(row=rr, column=5, value=f'=C{rr}+D{rr}')
            su.cell(row=rr, column=6, value=f'=IF($E${tot}=0,0,E{rr}/$E${tot})').number_format = '0%'
        return start + 2 + len(values) + 1

    nxt = breakdown(tot + 2, 'Cards per region', 'G', sorted({r['region'] for r in topics}))
    nxt = breakdown(nxt, 'Cards per strand', 'H', sorted({r['strand'] for r in topics}))
    nxt = breakdown(nxt, 'Cards per era', 'F', sorted({r['era'] for r in topics}))
    su.cell(row=nxt, column=1, value='Outcomes per level').font = Font(name=F, size=12, bold=True)
    for c, h in enumerate(['Level', 'Outcomes'], 1):
        cell = su.cell(row=nxt + 1, column=c, value=h)
        cell.font, cell.fill = head_font, head_fill
    lo_last = len(outs) + 1
    for j, lv in enumerate(LEVELS):
        rr = nxt + 2 + j
        su.cell(row=rr, column=1, value=f'Level {lv}')
        su.cell(row=rr, column=2, value=f"=COUNTIF('Level outcomes'!$A$2:$A${lo_last},\"{lv}\")")
    for row in su.iter_rows():
        for cell in row:
            if cell.font != head_font and not cell.font.bold:
                cell.font = base
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 110
    wb.save(path)


def main():
    when = sys.argv[sys.argv.index('--date') + 1] if '--date' in sys.argv else date.today().isoformat()
    m = json.loads(MAP.read_text(encoding='utf-8'))
    cards = json.loads((ROOT / 'data' / 'cards.json').read_text(encoding='utf-8'))
    problems = check(m, cards)
    for p in problems:
        print('PROBLEM:', p)
    if problems:
        sys.exit(f'{len(problems)} problem, ingen fil skriven')
    OUT.mkdir(exist_ok=True)
    path = OUT / f'Akelius-history-curriculum-map-{when}.xlsx'
    build(m, cards, path, when)
    per_tier = Counter()
    for r in m['topics']:
        per_tier[r['tier']] += r['existing_cards'] + r['new_cards']
    print(f'skrev {path.relative_to(ROOT)}: {len(m["topics"])} ämnesrader, {len(m["level_outcomes"])} mål, '
          f'kort per nivågrupp {dict(per_tier)}, totalt {sum(per_tier.values())}')


if __name__ == '__main__':
    main()
