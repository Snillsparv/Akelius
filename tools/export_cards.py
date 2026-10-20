#!/usr/bin/env python3
"""Exportera alla kort till granskningsbara filer för Akelius layoutteam.

Skapar i export/:
  Akelius-quiz-cards-EN-<datum>.pdf   stående A6, två sidor per kort, engelska
  Akelius-quiz-cards-SV-<datum>.pdf   samma på svenska
  Akelius-quiz-cards-text-<datum>.xlsx all text, en rad per kort, EN och SV

Sidlayouten följer Rogers memo 2026-10-15 (stående A6 för tryck och
smartphone). Framsidan uppifrån: huvudbilden 3:2 i full bredd, ledtrådstext,
fråga, fyra svar i 2x2. Baksidan: svaret med exakt samma ord som rätt
alternativ i Verdana 15, sedan ordförklaringarna som text, bild eller båda
enligt ordets visning. Ordlisteord är fetade i texten. Saknade bilder visas
som grå platshållare, huvudbilden med bildbriefen.

PDF:erna renderas med tools/export_pdf.js (Playwright + systemets Chromium).
Framsidans text sätts i Verdana 11 och stegas ned till 7 tills den ryms;
storleken står i sidfoten och i Excelfilen. Kör:
  python3 tools/export_cards.py [--date 2026-10-20] [--only slug1,slug2]
                                [--skip-pdf]
Med --only blir det ett urval: bara PDF:er, med filnamnsändelsen -sample.
"""
import html, json, math, pathlib, re, subprocess, sys
from datetime import date

ROOT = pathlib.Path(__file__).resolve().parent.parent
IMAGES = ROOT / 'images'
OUT = ROOT / 'export'
BUILD = OUT / 'build'
IMG_OUT = BUILD / 'img'
PAGES_URL = 'https://snillsparv.github.io/Akelius/images/'

MAIN_W, WORD_W, JPEG_Q = 760, 420, 72
TOC_ROWS = 30  # innehållsrader per A6-sida

T = {
    'en': {
        'pending': 'Picture in production', 'brief': 'Picture brief',
        'topic': 'Topic', 'card': 'card {k} of {n}', 'page': 'page',
        'levels': {'grade 6': 'Grade 6', 'grade 9': 'Grade 9',
                   'grade 12': 'Grade 12', 'university': 'University'},
        'contents': 'Contents', 'pics_pending': 'pictures in production',
        'cover_title': 'Akelius Quiz Cards',
        'cover_sub': 'World history',
        'cover_count': '{t} topics, {c} cards, English',
        'cover_date': 'Content export, {d}',
        'cover_format': 'A6 portrait, front and back',
        'cover_by': 'Text and pictures: Jonas von Essen',
        'legend_h': 'How to read this file',
        'legend': [
            'Each card has two pages, first the front and then the back. The '
            'pages are A6 portrait for printed cards and smartphones.',
            'Front: picture, clue text, question and four answers. Words in '
            'bold are explained on the back.',
            'Back: the answer in the exact words of the correct option, then '
            'the word explanations as text, picture or both.',
            'The front text is Verdana 11 where it fits and smaller where it '
            'does not. The footer shows the size.',
            'Every topic has five cards: grade 6, grade 9, two for grade 12 '
            'and university. Each card stands alone.',
            'All text has passed an independent historian review and an '
            'editor review. Every card also exists in Swedish.',
            'The pictures are AI-generated and reviewed one by one. A grey '
            'box marks a picture still in production.',
        ],
        'legend_pending': '{p} topics still have pictures in production; they are '
                          'marked with a star in the contents.',
    },
    'sv': {
        'pending': 'Bild under produktion', 'brief': 'Bildbrief på engelska',
        'topic': 'Ämne', 'card': 'kort {k} av {n}', 'page': 'sida',
        'levels': {'grade 6': 'Årskurs 6', 'grade 9': 'Årskurs 9',
                   'grade 12': 'Årskurs 12', 'university': 'Universitet'},
        'contents': 'Innehåll', 'pics_pending': 'bilder under produktion',
        'cover_title': 'Akelius frågekort',
        'cover_sub': 'Världshistoria',
        'cover_count': '{t} ämnen, {c} kort, svenska',
        'cover_date': 'Innehållsexport, {d}',
        'cover_format': 'Stående A6, framsida och baksida',
        'cover_by': 'Text och bild: Jonas von Essen',
        'legend_h': 'Så läser du filen',
        'legend': [
            'Varje kort har två sidor, först framsidan och sedan baksidan. '
            'Sidorna är stående A6 för tryckta kort och smartphone.',
            'Framsida: bild, ledtrådstext, fråga och fyra svar. Fetade ord '
            'förklaras på baksidan.',
            'Baksida: svaret med exakt samma ord som rätt alternativ, sedan '
            'ordförklaringarna som text, bild eller båda.',
            'Framsidans text är Verdana 11 där den ryms och mindre där den '
            'inte gör det. Sidfoten visar storleken.',
            'Varje ämne har fem kort: årskurs 6, årskurs 9, två för årskurs 12 '
            'och universitet. Varje kort fungerar fristående.',
            'All text har gått igenom oberoende historikergranskning och '
            'redaktörsgranskning. Varje kort finns också på engelska.',
            'Bilderna är AI-genererade och granskade en och en. En grå ruta '
            'markerar en bild under produktion.',
        ],
        'legend_pending': '{p} ämnen har fortfarande bilder under produktion; de är '
                          'markerade med stjärna i innehållet.',
    },
}

CAT_SV = {'civilizations': 'civilisation', 'religions': 'religion',
          'middle ages': 'medeltiden', 'renaissance': 'renässansen',
          'exploration': 'upptäcktsresor', 'science': 'vetenskap',
          'revolutions': 'revolutioner', 'ideas': 'idéhistoria',
          'twentieth century': '1900-talet', 'mathematics': 'matematik',
          'africa': 'Afrika', 'colonialism': 'kolonialism', 'asia': 'Asien',
          'latin america': 'Latinamerika', 'middle east': 'Mellanöstern'}

CAT_EN = {'africa': 'Africa', 'asia': 'Asia', 'latin america': 'Latin America',
          'middle east': 'Middle East'}

SUFFIX = {'en': r'(?:s|es|ed|d|ing|er|ers)?',
          'sv': r'(?:en|et|n|t|na|er|ar|or|erna|arna|orna|a|e|s|ande|de|te|r|ens|ets|s)?'}

CSS = """
@page { size: 105mm 148mm; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; }
body { font-family: Verdana, 'DejaVu Sans', sans-serif; color: #111; }
.page { width: 105mm; height: 148mm; padding: 6mm 6mm 5mm; position: relative;
  overflow: hidden; display: flex; flex-direction: column;
  page-break-after: always; break-after: page; }
.page:last-child { page-break-after: auto; break-after: auto; }
.foot { flex: none; display: flex; justify-content: space-between; gap: 3mm;
  font-size: 6.5pt; color: #555; padding-top: 1.5mm; margin-top: auto; }
.nw { white-space: nowrap; }
/* framsida */
.pic { width: 93mm; height: 62mm; flex: none; object-fit: cover; display: block; background: #eee; }
.content { flex: 1; min-height: 0; overflow: hidden; margin-top: 3mm; font-size: var(--fs, 11pt); }
.text { line-height: 1.3; }
.text p { margin: 0; display: inline; }
.q { font-weight: bold; margin-top: 1.6mm; line-height: 1.25; }
.opts { display: grid; grid-template-columns: 1fr 1fr; gap: 1.2mm 3mm; margin-top: 1.6mm; }
.opts div { display: flex; gap: 1.6mm; line-height: 1.2; }
.opts i { width: 3.2mm; height: 3.2mm; border: 0.3mm solid #111; flex: none; margin-top: 0.3mm; }
.ph { background: #eee; border: 0.3mm dashed #999; color: #444; font-size: 7pt;
  padding: 3mm; line-height: 1.35; overflow: hidden; }
.ph b { display: block; font-size: 7.5pt; margin-bottom: 1mm; color: #222; }
/* baksida */
.ans { flex: none; font-size: 15pt; font-weight: bold; line-height: 1.25; margin-bottom: 4mm; }
.expl { flex: 1; min-height: 0; overflow: hidden; font-size: var(--ws, 9.5pt); line-height: 1.3; }
.tiles { display: grid; gap: 3mm; margin-bottom: 3mm; }
.tiles.n2 { grid-template-columns: 1fr 1fr; }
.tiles.n3 { grid-template-columns: 1fr 1fr 1fr; }
.tile .img { width: 100%; aspect-ratio: 1; object-fit: cover; display: block; background: #eee; }
.tile .ph.img { padding: 2mm; font-size: 6.5pt; }
.tile .w { font-weight: bold; margin-top: 1.2mm; }
.tiles.n1 .tile { display: flex; gap: 4mm; align-items: flex-start; }
.tiles.n1 .tile .img { width: 46mm; flex: none; }
.tiles.n1 .tile .w { margin-top: 0; }
.lines div { margin-bottom: 1.8mm; }
/* omslag, läsanvisning och innehåll */
.cover h1 { font-size: 17pt; margin: 18mm 0 2mm; }
.cover .sub { font-size: 11pt; color: #333; margin-bottom: 8mm; }
.cover .meta { font-size: 8.5pt; line-height: 1.7; }
.legend h2 { font-size: 10pt; margin: 0 0 3mm; }
.legend ul { font-size: 7.6pt; line-height: 1.4; margin: 0; padding-left: 4mm; }
.legend li { margin-bottom: 1.6mm; }
.toc h2 { font-size: 10pt; margin: 0 0 3mm; }
.toc .row { display: flex; font-size: 7.4pt; line-height: 1.45; }
.toc .row .n { width: 7mm; color: #666; }
.toc .row .t { flex: 1; overflow: hidden; white-space: nowrap; text-overflow: ellipsis; }
.toc .row .p { width: 11mm; text-align: right; }
.toc .note { font-size: 6.5pt; color: #555; margin-top: 2mm; }
"""


def esc(s):
    return html.escape(str(s), quote=True)


def brief(s):
    """Bildbriefer är interna arbetsanteckningar med tankstreck; utåt blir de kommatecken."""
    return re.sub(r'\s*[—–]\s*|\s+-\s+', ', ', s or '')


def motif(s, limit=300):
    """Platshållarens text: briefens motiv utan produktionsnoten, kapat vid meningsgräns."""
    m = re.split(r'(?i)\s*(?:important\s+)?production note\s*:', brief(s))[0].strip()
    # meningar om vilket engelskt ordlisteord bilden förklarar hör inte hemma utåt
    m = re.sub(r"(?i)(?:^|(?<=[.;] ))(?:(?:it|this|the)(?:\s+(?:picture|image|photo))?\s+)?(?:also\s+)?"
               r"(?:helps?\s+(?:to\s+)?)?(?:explains?|supports?|illustrates?)\b[^.;]*\b(?:words?|phrase)\b"
               r"[^.;]*[.;]\s*", '', m).strip()
    if len(m) <= limit:
        return m
    cut = max(m.rfind('. ', 0, limit), m.rfind('; ', 0, limit))
    return (m[:cut + 1] if cut > limit // 2 else m[:m.rfind(' ', 0, limit)].rstrip(',;')) + ' …'


def image_ids():
    return {p.stem for p in IMAGES.glob('*.png')}


def used_images(sets, have):
    """Bilder som syns i A6-layouten: huvudbilder och ordbilder. Sidobilden har ingen plats."""
    ids = set()
    for s in sets:
        for k, card in enumerate(s['cards'], 1):
            ids.add(f"{s['slug']}-{k}-main")
            ids.update(w['picture']['id'] for w in card['words'] if 'picture' in w)
    return ids & have


def prepare_images(ids):
    from PIL import Image
    IMG_OUT.mkdir(parents=True, exist_ok=True)
    stamp = IMG_OUT / '.params'
    params = f'{MAIN_W}-{WORD_W}-{JPEG_Q}'
    fresh = stamp.exists() and stamp.read_text() == params
    made = 0
    for iid in sorted(ids):
        src = IMAGES / f'{iid}.png'
        dst = IMG_OUT / f'{iid}.jpg'
        if fresh and dst.exists() and dst.stat().st_mtime >= src.stat().st_mtime:
            continue
        w = WORD_W if '-w-' in iid else MAIN_W
        im = Image.open(src).convert('RGB')
        h = round(im.height * w / im.width)
        im.resize((w, h), Image.LANCZOS).save(dst, 'JPEG', quality=JPEG_Q,
                                               optimize=True, progressive=True)
        made += 1
    stamp.write_text(params)
    return made


def topic_name(s, lang):
    """Visningsnamn för ett ämne; tankstreck i namnet blir kolon."""
    name = s.get('person_sv', s['person']) if lang == 'sv' else s['person']
    return name.replace(' - ', ': ')


STEM_RULES = {
    'en': (('man', 'men'), ('y', ''), ('e', '')),
    # handelsman/handelsmän, kloster/klostret, partikel/partiklar, öken/öknen
    'sv': (('man', 'män'), ('are', 'ar'), ('ium', 'i'), ('a', ''), ('e', ''),
           ('er', 'r'), ('el', 'l'), ('en', 'n')),
}


def stems(word, lang, minlen=4):
    """Ordstammar som tål böjningar där ändelsen byts: källa/källor, century/centuries."""
    out = [word]
    for end, new in STEM_RULES[lang]:
        if word.lower().endswith(end):
            out.append(word[:-len(end)] + new)
    return [s for s in out if len(s) >= minlen]


def term_patterns(term, lang):
    """Mönster i fallande säkerhet: exakt ord med ändelse, stam som också fångar
    sammansättningar (laser i laserljus), och för partikelverb (köra om, hitta på)
    verbet med partikeln, eller bara verbet när partikeln följer några ord senare."""
    words = term.split()
    yield r'(?<!\w)(' + re.escape(esc(term)) + SUFFIX[lang] + r')(?!\w)'
    if len(words) == 1:
        for st in stems(term, lang):
            yield r'(?<!\w)(' + re.escape(esc(st)) + r'\w*)'
        return
    first = [re.escape(esc(st)) for st in stems(words[0], lang, minlen=3) or [words[0]]]
    rest = r'\s+'.join(re.escape(esc(x)) for x in words[1:]) + r'(?!\w)'
    for st in first:
        yield r'(?<!\w)(' + st + r'\w*\s+' + rest + ')'
    for st in first:
        yield r'(?<!\w)(' + st + r'\w*)(?=(?:\s+[^\s<]+){0,4}\s+' + rest + ')'


def bold_words(sentences, words, lang):
    """Feta första förekomsten av varje ordlisteord, som i Akelius skiss. Förekomster
    med samma skiftläge som ordlistan går före, så att "stat" inte fetas i "Staten"."""
    out = [esc(s) for s in sentences]

    def bold_first(pat):
        for i, s in enumerate(out):
            # hoppa över redan fetad text för att inte nästla taggar
            parts = re.split(r'(<b>.*?</b>)', s)
            for j, part in enumerate(parts):
                if part.startswith('<b>'):
                    continue
                new, n = pat.subn(r'<b>\1</b>', part, count=1)
                if n:
                    parts[j] = new
                    out[i] = ''.join(parts)
                    return True
        return False

    for w in words:
        term = w['word'].strip()
        if term:
            any(bold_first(re.compile(p, flags))
                for p in term_patterns(term, lang) for flags in (0, re.IGNORECASE))
    # 1430-talet och liknande bryts aldrig vid bindestrecket
    return [re.sub(r'(\w+(?:-\w+)+)', r'<span class="nw">\1</span>', x) for x in out]


def localized(card, lang):
    """Kortets text på valt språk. Ordens visning och ordbild följer det engelska
    ordet på samma plats i listan."""
    if lang == 'en':
        L = {
            'title': card['title'], 'sentences': card['text_sentences'],
            'question': card['question'], 'options': card['options'],
            'correct': card['correct'], 'words': card['words'],
            'main_cap': card['main_image']['caption'],
            'side_cap': card['side_image']['caption'],
            'category': CAT_EN.get(card['category'], card['category']),
        }
    else:
        sv = card['sv']
        L = {
            'title': sv['title'], 'sentences': sv['text_sentences'],
            'question': sv['question'], 'options': sv['options'],
            'correct': sv['correct'],
            'words': [dict(w, show=e.get('show', 'text'), **({'picture': e['picture']}
                                                              if 'picture' in e else {}))
                      for w, e in zip(sv['words'], card['words'])],
            'main_cap': sv['main_caption'], 'side_cap': sv['side_caption'],
            'category': CAT_SV.get(card['category'], card['category']),
        }
    return L


def main_picture(iid, have, desc, t):
    if iid in have:
        return f'<img class="pic" src="img/{iid}.jpg" alt="">'
    return (f'<div class="pic ph"><b>{esc(t["pending"])}</b>'
            f'{esc(t["brief"])}: {esc(motif(desc))}</div>')


def word_picture(w, have, t):
    pid = w['picture']['id']
    if pid in have:
        return f'<img class="img" src="img/{pid}.jpg" alt="">'
    return f'<div class="img ph">{esc(t["pending"])}</div>'


def back_words(words, have, t):
    """Ordförklaringarna: ord med bild som rutor överst, rena textord som rader under."""
    def pictured(w):
        return w.get('show', 'text') != 'text' and 'picture' in w
    tiles = [w for w in words if pictured(w)]
    lines = [w for w in words if not pictured(w)]
    out = ''
    if tiles:
        cells = ''.join(
            f'<div class="tile">{word_picture(w, have, t)}<div><div class="w">{esc(w["word"])}</div>'
            + (f'<div class="e">{esc(w["explanation"])}</div>' if w['show'] == 'both' else '')
            + '</div></div>' for w in tiles)
        out += f'<div class="tiles n{min(len(tiles), 3)}">{cells}</div>'
    if lines:
        out += '<div class="lines">' + ''.join(
            f'<div><b>{esc(w["word"])}</b>: {esc(w["explanation"])}</div>' for w in lines) + '</div>'
    return out


def card_pages(set_no, s, k, card, have, lang, page_no):
    t = T[lang]
    L = localized(card, lang)
    topic = topic_name(s, lang)
    iid = f"{s['slug']}-{k}"
    cid = f'{set_no:03d}-{k}'
    sentences = bold_words(L['sentences'], L['words'], lang)
    opts = ''.join(f'<div><i></i><span>{esc(o)}</span></div>' for o in L['options'])
    level = t['levels'].get(card['level'], card['level'])
    question = re.sub(r' X(?!\w)', ' X', L['question'])  # "X?" hamnar aldrig ensamt på en rad
    card_of = t['card'].format(k=k, n=len(s['cards']))
    return f"""
<section class="page card front" data-id="{iid}">
  {main_picture(iid + '-main', have, card['main_image']['description'], t)}
  <div class="content">
    <div class="text">{''.join(f'<p>{x} </p>' for x in sentences)}</div>
    <div class="q">{esc(question)}</div>
    <div class="opts">{opts}</div>
  </div>
  <div class="foot"><span>{esc(level)} · {esc(L['category'])}</span>
    <span>{cid} · <span class="fsz">11</span> pt · {esc(t['page'])} {page_no}</span></div>
</section>
<section class="page card back" data-id="{iid}">
  <div class="ans">{esc(L['correct'])}</div>
  <div class="expl">{back_words(L['words'], have, t)}</div>
  <div class="foot"><span>{esc(t['topic'])} {set_no}: {esc(topic)}, {esc(card_of)}</span>
    <span>{cid} · {esc(t['page'])} {page_no + 1}</span></div>
</section>"""


def pending_sets(sets, have):
    """Ämnen där någon bild som syns i A6-layouten saknas."""
    return [s for s in sets if any(
        f"{s['slug']}-{k}-main" not in have
        or any(w['picture']['id'] not in have for w in c['words'] if 'picture' in w)
        for k, c in enumerate(s['cards'], 1))]


def build_html(data, have, lang, when, numbers):
    """numbers: ämnets nummer i hela samlingen, så att ett urval behåller kort-ID:n."""
    t = T[lang]
    sets = data['sets']
    pending = pending_sets(sets, have)
    pend_slugs = {s['slug'] for s in pending}
    n_cards = sum(len(s['cards']) for s in sets)
    toc_pages = math.ceil(len(sets) / TOC_ROWS)
    legend = t['legend'] + ([t['legend_pending'].format(p=len(pending))] if pending else [])
    first_card_page = 3 + toc_pages
    pages = [f"""
<section class="page cover">
  <h1>{esc(t['cover_title'])}</h1>
  <div class="sub">{esc(t['cover_sub'])}</div>
  <div class="meta">{esc(t['cover_count'].format(t=len(sets), c=n_cards))}<br>
    {esc(t['cover_format'])}<br>{esc(t['cover_date'].format(d=when))}<br>{esc(t['cover_by'])}</div>
  <div class="foot"><span></span><span>{esc(t['page'])} 1</span></div>
</section>""", f"""
<section class="page legend">
  <h2>{esc(t['legend_h'])}</h2>
  <ul>{''.join(f'<li>{esc(x)}</li>' for x in legend)}</ul>
  <div class="foot"><span></span><span>{esc(t['page'])} 2</span></div>
</section>"""]
    rows, p = [], first_card_page
    for s in sets:
        star = ' *' if s['slug'] in pend_slugs else ''
        rows.append(f'<div class="row"><span class="n">{numbers[s["slug"]]}</span>'
                    f'<span class="t">{esc(topic_name(s, lang))}{star}</span>'
                    f'<span class="p">{p}</span></div>')
        p += 2 * len(s['cards'])
    for i in range(toc_pages):
        note = f'<div class="note">* {esc(t["pics_pending"])}</div>' if pend_slugs else ''
        pages.append(f"""
<section class="page toc">
  <h2>{esc(t['contents'])}</h2>
  {''.join(rows[i * TOC_ROWS:(i + 1) * TOC_ROWS])}
  {note}
  <div class="foot"><span></span><span>{esc(t['page'])} {3 + i}</span></div>
</section>""")
    p = first_card_page
    for s in sets:
        for k, card in enumerate(s['cards'], 1):
            pages.append(card_pages(numbers[s['slug']], s, k, card, have, lang, p))
            p += 2
    doc = (f'<!DOCTYPE html><html lang="{lang}"><head><meta charset="utf-8">'
           f'<title>{esc(t["cover_title"])}</title><style>{CSS}</style></head>'
           f'<body>{"".join(pages)}</body></html>')
    return doc, p - 1, len(pending)


def build_xlsx(data, have, path, when, fits):
    """fits: {språk: {kort-id: {'front': pt, 'back': pt}}} från PDF-renderingen."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    readme = wb.active
    readme.title = 'Read me'
    n_cards = sum(len(s['cards']) for s in data['sets'])
    lines = [
        ('Akelius Quiz Cards · World history', True),
        (f'Content export {when}. {len(data["sets"])} topics, {n_cards} cards, '
         'English and Swedish. Layout: A6 portrait, front and back.', False),
        ('', False),
        ('Sheets', True),
        ('Topics: one row per topic, with picture status.', False),
        ('Cards EN and Cards SV: one row per card. Text has one sentence per line.', False),
        ('', False),
        ('Columns', True),
        ('Card ID: topic number and card number, for example 012-3.', False),
        ('Front text size: the largest Verdana size in points where text, question and '
         'answers fit on the A6 front below the picture. 11 is the house size.', False),
        ('Correct: the right answer, also given as a letter A to D in the order of the options. '
         'The back of the card repeats these exact words.', False),
        ('Words: word explanations, one per line as word: explanation.', False),
        ('Word display: how each word is shown on the back, text, picture or both.', False),
        ('Word pictures: file name and link for each word shown with a picture.', False),
        ('Main picture: file name, link to the full-resolution file, caption, '
         'picture brief and a Shutterstock search query for replacing the AI picture with a photo.', False),
        ('Side picture: the second picture of the earlier landscape layout. It has no place '
         'in the A6 portrait layout and is kept here until its role is decided.', False),
        ('', False),
        ('Notes', True),
        ('Picture file names contain the topic name, which is the answer. Rename them '
         'before any student-facing use.', False),
        ('Pictures are AI-generated and reviewed one by one. Where the link is empty, '
         'the picture is still in production.', False),
        ('Card text, title, captions and word list never name the answer.', False),
    ]
    for r, (txt, bold) in enumerate(lines, 1):
        c = readme.cell(row=r, column=1, value=txt)
        c.font = Font(bold=bold, size=12 if bold else 11)
    readme.column_dimensions['A'].width = 120

    topics = wb.create_sheet('Topics')
    topics.append(['#', 'Topic (EN)', 'Topic (SV)', 'Category', 'Cards',
                   'Main pictures ready', 'Word pictures ready', 'Slug'])
    for i, s in enumerate(data['sets'], 1):
        n = len(s['cards'])
        main = sum(f"{s['slug']}-{k}-main" in have for k in range(1, n + 1))
        wids = [w['picture']['id'] for c in s['cards'] for w in c['words'] if 'picture' in w]
        topics.append([i, topic_name(s, 'en'), topic_name(s, 'sv'), s['cards'][0]['category'],
                       n, f'{main} of {n}', f'{sum(w in have for w in wids)} of {len(wids)}',
                       s['slug']])

    show_name = {'text': 'text', 'picture': 'picture', 'both': 'text and picture'}

    def card_rows(lang):
        rows = []
        for i, s in enumerate(data['sets'], 1):
            for k, card in enumerate(s['cards'], 1):
                L = localized(card, lang)
                iid = f"{s['slug']}-{k}"
                letter = 'ABCD'[L['options'].index(L['correct'])]
                level = T[lang]['levels'].get(card['level'], card['level'])
                fit = fits.get(lang, {}).get(iid, {}).get('front', '')
                pics = [w for w in L['words'] if 'picture' in w]
                row = [f'{i:03d}-{k}', i, topic_name(s, lang), k, level, L['category'], L['title'],
                       '\n'.join(L['sentences']), fit, L['question'], *L['options'],
                       L['correct'], letter,
                       '\n'.join(f"{w['word']}: {w['explanation']}" for w in L['words']),
                       '\n'.join(f"{w['word']}: {show_name[w.get('show', 'text')]}" for w in L['words']),
                       '\n'.join(f"{w['word']}: {w['picture']['id']}.png"
                                 + (f" {PAGES_URL}{w['picture']['id']}.png"
                                    if w['picture']['id'] in have else ' (in production)')
                                 for w in pics)]
                for kind, cap in (('main', L['main_cap']), ('side', L['side_cap'])):
                    pid = f'{iid}-{kind}'
                    img = card[f'{kind}_image']
                    row += [f'{pid}.png', PAGES_URL + f'{pid}.png' if pid in have else '',
                            cap, brief(img['description']), brief(img['shutterstock_query'])]
                rows.append(row)
        return rows

    cols = ['Card ID', 'Topic #', 'Topic', 'Card #', 'Level', 'Category', 'Title',
            'Text', 'Front text size (pt)', 'Question', 'Option A', 'Option B', 'Option C',
            'Option D', 'Correct', 'Correct letter', 'Words', 'Word display', 'Word pictures',
            'Main picture file', 'Main picture link', 'Main caption', 'Main picture brief',
            'Main Shutterstock query',
            'Side picture file', 'Side picture link', 'Side caption', 'Side picture brief',
            'Side Shutterstock query']
    widths = [9, 7, 26, 6, 11, 14, 30, 70, 9, 22, 18, 18, 18, 18, 18, 7, 45, 28, 50,
              26, 30, 36, 45, 30, 26, 30, 36, 45, 30]
    for name, lang in (('Cards EN', 'en'), ('Cards SV', 'sv')):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for row in card_rows(lang):
            ws.append(row)
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
    for ws in (topics, wb['Cards EN'], wb['Cards SV']):
        for cell in ws[1]:
            cell.font = Font(bold=True)
    topics.freeze_panes = 'A2'
    for c, w in enumerate([5, 40, 40, 18, 7, 14, 14, 28], 1):
        topics.column_dimensions[get_column_letter(c)].width = w
    wb.save(path)


def main():
    args = sys.argv[1:]
    when = args[args.index('--date') + 1] if '--date' in args else date.today().isoformat()
    only = args[args.index('--only') + 1].split(',') if '--only' in args else None
    data = json.loads((ROOT / 'data' / 'cards.json').read_text(encoding='utf-8'))
    numbers = {s['slug']: i for i, s in enumerate(data['sets'], 1)}
    if only:
        unknown = set(only) - set(numbers)
        if unknown:
            sys.exit(f'Okända ämnen: {sorted(unknown)}')
        data = dict(data, sets=[s for s in data['sets'] if s['slug'] in only])
    missing_sv = [s['slug'] for s in data['sets'] if not all('sv' in c for c in s['cards'])]
    if missing_sv:
        sys.exit(f'Svenska saknas för: {missing_sv}')
    have = image_ids()
    BUILD.mkdir(parents=True, exist_ok=True)
    made = prepare_images(used_images(data['sets'], have))
    print(f'bilder: {len(have)} finns, {made} nya exportkopior')
    tail = f'{when}-sample' if only else when
    fits = {}
    for lang in ('en', 'sv'):
        doc, last_page, n_pending = build_html(data, have, lang, when, numbers)
        html_path = BUILD / f'cards-{lang}.html'
        html_path.write_text(doc, encoding='utf-8')
        pdf = OUT / f'Akelius-quiz-cards-{lang.upper()}-{tail}.pdf'
        print(f'{lang}: {last_page} sidor, {n_pending} ämnen med bilder under produktion')
        if '--skip-pdf' in args:
            continue
        report = BUILD / f'fit-{lang}.json'
        subprocess.run(['node', str(ROOT / 'tools' / 'export_pdf.js'), str(html_path),
                        str(pdf), str(last_page), str(report)], check=True)
        fits[lang] = json.loads(report.read_text(encoding='utf-8'))
    if not only:
        xlsx = OUT / f'Akelius-quiz-cards-text-{when}.xlsx'
        build_xlsx(data, have, xlsx, when, fits)
        print('skrev', xlsx.relative_to(ROOT))


if __name__ == '__main__':
    main()
